"""P33a step 3 (B151 §3c): a paired response workbook renders as the
markdown the ingest reads — one chunk per answered question, the question
as the card's title, the answer as its body. Expectations are hand-derived
from the in-test workbooks and the committed demo twin, never pasted.
"""

import json
from pathlib import Path

from openpyxl import Workbook

from engine.kb import KBStore, SourceDoc, ingest_document
from engine.kb.canonical import elements_from_markdown
from engine.kb.chunk import chunk_elements
from engine.kb.response_workbook import (
    RenderedMarkdown,
    is_markdown_structural,
    pair_response_workbook,
    render_markdown,
)
from engine.llm import FakeCaller, TracedCaller
from engine.runlog import RunLogger
from tests.fixtures.filled_twins import fill_demo_twin

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
BUYER = "Northwind Regional Health"
WIRE = json.dumps({"chunk_annotations": [], "qa_pairs": [], "identifiers": [],
                   "client_descriptor": "a regional health system"})
REVIEW = json.dumps({"identifiers": []})


def _sections_workbook(path: Path, rows=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2. Delivery "  # the trailing space is real (EC-1)
    ws["A1"], ws["B1"], ws["C1"] = "Ref", "Question", "Response"
    rows = rows or [
        ("1", "Company background", None),
        ("1.1", "Describe your company history.",
         "Founded in 1998 and privately held, with offices in four regions."),
        ("1.2", f"Describe your experience with {BUYER}.",
         "Two prior programs, both delivered on the agreed timeline."),
        ("2", "Delivery approach", None),
        ("2.1", "Describe your methodology.",
         "Phased delivery with gates.\n\nNamed owners and weekly steering."),
    ]
    for i, (ref, q, a) in enumerate(rows, start=2):
        ws[f"A{i}"], ws[f"B{i}"] = ref, q
        if a:
            ws[f"C{i}"] = a
    wb.save(path)
    return path


def _chunks(text: str):
    elements = elements_from_markdown(text)
    return elements, chunk_elements(elements)


def _body(elements, chunk) -> str:
    start, end = chunk.elements
    return "\n".join(e.text for e in elements[start:end])


def test_rendered_markdown_chunks_one_per_question(tmp_path):
    book = pair_response_workbook(_sections_workbook(tmp_path / "sec.xlsx"))
    out = render_markdown(book)
    assert isinstance(out, RenderedMarkdown)
    assert out.rendered == 3 and out.warnings == []
    elements, chunks = _chunks(out.text)
    assert [c.doc_path for c in chunks] == [
        ["2. Delivery", "Company background", "Describe your company history."],
        ["2. Delivery", "Company background", f"Describe your experience with {BUYER}."],
        ["2. Delivery", "Delivery approach", "Describe your methodology."],
    ]
    assert _body(elements, chunks[0]) == (
        "Founded in 1998 and privately held, with offices in four regions.")
    # a blank line inside an answer splits paragraphs, never the chunk
    assert _body(elements, chunks[2]) == (
        "Phased delivery with gates.\nNamed owners and weekly steering.")
    assert chunks[2].elements[1] - chunks[2].elements[0] == 2


def test_sheet_without_sections_paths_sheet_then_question(tmp_path):
    book = pair_response_workbook(fill_demo_twin(tmp_path / "filled.xlsx"))
    out = render_markdown(book)
    assert out.rendered == 19
    _elements, chunks = _chunks(out.text)
    assert len(chunks) == 19
    assert all(len(c.doc_path) == 2 for c in chunks)
    assert chunks[0].doc_path == [
        "1. Implementation Methodology",
        "Describe your implementation methodology for ERP deployments."]
    assert not any(c.doc_path == ["1. Implementation Methodology"] for c in chunks)


def test_question_whitespace_is_collapsed_so_it_never_reaches_the_body(tmp_path):
    rows = [("1.1", "Describe your approach\nto data migration.",
             "Phased, rehearsed, and reconciled before cutover.")]
    book = pair_response_workbook(_sections_workbook(tmp_path / "nl.xlsx", rows))
    out = render_markdown(book)
    assert "### Describe your approach to data migration." in out.text.splitlines()
    elements, chunks = _chunks(out.text)
    assert len(chunks) == 1
    assert _body(elements, chunks[0]) == "Phased, rehearsed, and reconciled before cutover."


def test_pipe_in_question_stays_a_heading(tmp_path):
    rows = [("1.1", "Describe HL7 | FHIR interface support.",
             "Both, through the integration engine's standard adapters.")]
    book = pair_response_workbook(_sections_workbook(tmp_path / "pipe.xlsx", rows))
    elements, chunks = _chunks(render_markdown(book).text)
    assert [e.kind for e in elements if "HL7" in e.text] == ["heading"]
    assert chunks[0].doc_path[-1] == "Describe HL7 | FHIR interface support."


def test_markdown_structural_answer_lines_are_refused(tmp_path):
    rows = [("1.1", "How many FTEs do you propose?", "# of FTEs: 40\nAcross two teams."),
            ("1.2", "Summarise your rate card structure.", "| role | rate |"),
            ("1.3", "Any internal notes?", "<!-- do not share -->"),
            ("1.4", "Describe your methodology.",
             "Phased delivery with gates and named owners per workstream.")]
    book = pair_response_workbook(_sections_workbook(tmp_path / "structural.xlsx", rows))
    out = render_markdown(book)
    assert out.rendered == 1
    assert out.warnings == [
        "2. Delivery!row 2: answer skipped (markdown-structural line)",
        "2. Delivery!row 3: answer skipped (markdown-structural line)",
        "2. Delivery!row 4: answer skipped (markdown-structural line)",
    ]
    assert "FTEs" not in out.text and "rate card" not in out.text.split("###")[0]
    _elements, chunks = _chunks(out.text)
    assert [c.doc_path for c in chunks] == [["2. Delivery", "Describe your methodology."]]
    assert is_markdown_structural("  ## heading  ") and not is_markdown_structural("#tag")


def test_provenance_comment_carries_no_filename(tmp_path):
    path = _sections_workbook(tmp_path / f"{BUYER} response 2025.xlsx")
    book = pair_response_workbook(path)
    out = render_markdown(book)
    first = out.text.splitlines()[0]
    assert first == (f"<!-- response-workbook sha256:{book.source_sha256[:12]} "
                     f"parser:{book.parser_version} -->")
    assert "response 2025" not in out.text and str(tmp_path) not in out.text
    elements, _chunks_ = _chunks(out.text)
    assert not any("sha256" in e.text for e in elements)  # dropped by the reader


def test_render_is_byte_deterministic(tmp_path):
    one = render_markdown(pair_response_workbook(fill_demo_twin(tmp_path / "a.xlsx")))
    two = render_markdown(pair_response_workbook(fill_demo_twin(tmp_path / "b.xlsx")))
    assert one.text == two.text
    assert one.text.endswith("\n") and not one.text.endswith("\n\n")


def test_rendered_markdown_ingests_one_card_per_question(tmp_path):
    book = pair_response_workbook(_sections_workbook(tmp_path / "sec.xlsx"))
    text = render_markdown(book).text
    store = KBStore(tmp_path / "kb")
    log = RunLogger(store.root, "run_0001", "kb")
    caller = TracedCaller(FakeCaller({"ingestion_agent": WIRE,
                                      "anonymization_reviewer": REVIEW}), log)
    doc = SourceDoc(doc_id="resp_sec", text=text, source_client=BUYER,
                    source_pursuit="pur_sec", outcome="won", date="2026-01-01",
                    authored_by="firm")
    report = ingest_document(store, caller, log, doc)
    assert report.status == "ingested", report.findings
    cards = [store.read_card(kb_id) for kb_id in report.cards_written]
    titles = sorted(card["title"] for card, _body_ in cards)
    assert titles == sorted([
        "Describe your company history.",
        "Describe your experience with [CLIENT].",  # the buyer, placeholdered
        "Describe your methodology.",
    ])
    bodies = {card["title"]: body.strip() for card, body in cards}
    assert bodies["Describe your company history."] == (
        "Founded in 1998 and privately held, with offices in four regions.")
    assert bodies["Describe your methodology."] == (
        "Phased delivery with gates.\nNamed owners and weekly steering.")
    assert all(BUYER not in title + body for title, body in bodies.items())
    assert all(card["doc_path"][-1] == card["title"] for card, _b in cards)
