"""P33a (B151): a completed response workbook pairs each buyer question with
the firm's answer beside it. Expectations are hand-derived from the
fixture builders (`tests/fixtures/twins.py`, `tests/fixtures/intake_twins.py`)
and the in-test workbooks, never pasted from a run.

Step 1 (P3-25): Layer 1 records hidden state and still collects the cells.
Step 2: the pairing (engine/structure/response.py, RESPONSE_PARSER_VERSION).
"""

import json
import re
from pathlib import Path

from openpyxl import Workbook

from engine.structure import PARSER_VERSION, parse_workbook
from engine.structure.conventions import learn_conventions
from engine.structure.facts import collect_workbook_facts
from engine.structure.response import (
    RESPONSE_PARSER_VERSION,
    ResponseWorkbook,
    ResponseWorkbookError,
    pair_response_workbook,
)
from tests.fixtures.filled_twins import default_answer, demo_rows, fill_demo_twin
from tests.fixtures.twins import _splice_member

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


def test_hidden_sheets_and_rows_are_recorded_at_layer_1():
    # Hand-derived from build_hidden_twin: "Vendor Questions" hides row 4
    # and column C; "Internal Notes" is a hidden sheet. Layer 1 RECORDS
    # the state and keeps every cell — hidden cells are still facts.
    facts = collect_workbook_facts(FIXTURES / "hidden-twin.xlsx")
    by_name = {s.name: s for s in facts.sheets}
    questions, notes = by_name["Vendor Questions"], by_name["Internal Notes"]
    assert notes.hidden is True
    assert questions.hidden is False
    assert questions.hidden_rows == {4}
    assert questions.hidden_cols == {3}
    assert notes.hidden_rows == set() and notes.hidden_cols == set()
    assert "B4" in questions.cells  # the hidden row's text is still a fact
    assert "C2" in questions.cells  # the hidden column's text too
    assert "A1" in notes.cells  # and the hidden sheet's


# ---------------------------------------------------------------------------
# Step 2: the pairing (engine/structure/response.py, RESPONSE_PARSER_VERSION).
# ---------------------------------------------------------------------------

_ROW_LINE = re.compile(r"^.+!row \d+: [a-z][a-z -]*[a-z]( \(.+\))?$")
_SHEET_LINE = re.compile(r"^.+: [a-z][a-z/, -]*[a-z]$")


def _shaped(line: str) -> bool:
    return bool(_ROW_LINE.match(line) or _SHEET_LINE.match(line))


def test_fully_filled_demo_twin_pairs_every_answer(tmp_path):
    # Hand-derived from build_demo_twin: 8 question sheets, 19 question
    # rows, one Instructions sheet. Filled, every row is a pair.
    book = pair_response_workbook(fill_demo_twin(tmp_path / "filled.xlsx"))
    assert isinstance(book, ResponseWorkbook)
    assert book.file == "filled.xlsx"
    assert book.parser_version == RESPONSE_PARSER_VERSION == "1.1.0"
    assert book.warnings == []
    assert book.skipped == {"instructions sheets": 1}
    assert len(book.sheets) == 8
    got = [(s.name, p.row, p.ref, p.question, p.answer)
           for s in book.sheets for p in s.pairs()]
    want = [(sheet, row, ref, question, default_answer(ref, question))
            for sheet, row, ref, question in demo_rows()]
    assert got == want
    assert len(got) == 19
    # the vacancy parser is untouched: same workbook, same pinned version
    assert parse_workbook(tmp_path / "filled.xlsx").parser_version == PARSER_VERSION


def test_half_filled_demo_twin_warns_unanswered_by_address(tmp_path):
    path = fill_demo_twin(tmp_path / "half.xlsx", rows="first")
    # the fill-mode leg is non-vacuous: the empty answer cells still vote
    assert learn_conventions(collect_workbook_facts(path)).writable_fill == "FFFF00"
    book = pair_response_workbook(path)
    assert len(book.pairs()) == 8  # one per sheet
    assert book.skipped == {"instructions sheets": 1, "unanswered": 11}
    # hand-listed: every row 3+ of every sheet, in sheet-then-row order
    assert book.warnings == [
        f"{sheet}!row {row}: unanswered"
        for sheet, row, _ref, _q in demo_rows() if row != 2
    ]
    assert all(_shaped(line) for line in book.warnings)


def test_short_question_short_answer_rows_pair_despite_being_label_rows(tmp_path):
    # Four demo questions are <= 60 chars (3.0.1, 3.0.2, 5.0.2, 8.0.1). A short
    # answer beside them makes the row a LABEL ROW to Layer 2 (two short
    # lettered texts, nothing longer) — the pairer must not skip them.
    short = {"3.0.1", "3.0.2", "5.0.2", "8.0.1"}
    path = fill_demo_twin(
        tmp_path / "short.xlsx",
        answers=lambda ref, q: "Yes, via HL7 adapters." if ref in short
        else default_answer(ref, q))
    sc = learn_conventions(collect_workbook_facts(path)).sheets["3. Integration"]
    assert {2, 3} <= set(sc.label_rows)  # the trap exists
    book = pair_response_workbook(path)
    assert len(book.pairs()) == 19
    assert book.warnings == []
    integration = next(s for s in book.sheets if s.name == "3. Integration")
    assert [p.answer for p in integration.pairs()] == ["Yes, via HL7 adapters."] * 2


def test_refless_two_column_sheet_pairs_by_header_not_kind(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws["A1"], ws["B1"] = "Question", "Response"
    ws["A2"], ws["B2"] = ("Years in business",
                          "We have operated continuously since 1998 across four regions.")
    ws["A3"], ws["B3"] = ("Describe your implementation methodology.",
                          "A phased approach with governance gates at each milestone.")
    ws["A4"], ws["B4"] = ("Describe your support model after go-live.",
                          "Hypercare for eight weeks, then a tiered service desk.")
    path = tmp_path / "refless.xlsx"
    wb.save(path)
    sc = learn_conventions(collect_workbook_facts(path)).sheets["Questions"]
    assert sc.kind == "grid"  # Layer 2 calls a filled ref-less sheet a grid
    assert sc.question_col == sc.answer_col  # and its vote has flipped
    book = pair_response_workbook(path)
    assert [(p.question, p.ref) for p in book.pairs()] == [
        ("Years in business", None),
        ("Describe your implementation methodology.", None),
        ("Describe your support model after go-live.", None),
    ]
    assert book.pairs()[0].answer.startswith("We have operated")
    assert book.warnings == []


def test_shallower_refs_without_answers_become_section_headings(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sec"
    # a long header label: found without a length cap
    ws["A1"], ws["B1"], ws["C1"] = (
        "Ref", "Requirement",
        "Vendor Response (max 200 words; do not alter formatting)")
    rows = [
        ("1", "Company background", None),
        ("1.1", "Describe your company history and ownership structure.",
         "Founded in 1998, privately held, offices in four regions."),
        ("1.2", "Describe your financial stability over the last three years.",
         "Audited statements show steady growth and no debt covenants."),
        ("2", "Delivery approach", None),
        ("2.1", "Describe your methodology.",
         "Phased delivery with gates, named owners and weekly steering."),
        ("3", "Any other remarks about your firm?",
         "We maintain partner certifications and a public-sector practice."),
    ]
    for i, (ref, q, a) in enumerate(rows, start=2):
        ws[f"A{i}"], ws[f"B{i}"] = ref, q
        if a:
            ws[f"C{i}"] = a
    path = tmp_path / "sections.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.warnings == []  # the header was FOUND, not inferred
    (sheet,) = book.sheets
    assert [s.heading for s in sheet.sections] == ["Company background",
                                                   "Delivery approach"]
    assert [p.row for p in sheet.sections[0].pairs] == [3, 4]
    # a shallow ref WITH an answer is a pair, never a heading
    assert [p.row for p in sheet.sections[1].pairs] == [6, 7]
    assert sheet.sections[1].pairs[1].ref == "3"


def test_fully_filled_one_dot_sheet_in_a_fill_mode_workbook_keeps_its_leaves(tmp_path):
    wb = Workbook()
    filled = wb.active
    filled.title = "A"
    filled["A1"], filled["B1"], filled["C1"] = "Ref", "Question", "Response"
    for i, ref in enumerate(("1.1", "1.2", "1.3"), start=2):
        filled[f"A{i}"] = ref
        filled[f"B{i}"] = f"Describe requirement {ref} in detail for the evaluation panel."
        filled[f"C{i}"] = f"Our answer to {ref}: a governed approach with named owners."
    empty = wb.create_sheet("B")
    empty["A1"], empty["B1"], empty["C1"] = "Ref", "Question", "Response"
    from openpyxl.styles import PatternFill
    fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    for i, ref in enumerate(("2.0.1", "2.0.2"), start=2):
        empty[f"A{i}"] = ref
        empty[f"B{i}"] = f"Describe requirement {ref} in detail for the evaluation panel."
        empty[f"C{i}"].fill = fill
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    conv = learn_conventions(collect_workbook_facts(path))
    assert conv.writable_fill is not None  # fill mode, learned from sheet B
    assert conv.sheets["A"].leaf_depth == 3  # the dataclass default — the trap
    book = pair_response_workbook(path)
    a_sheet = next(s for s in book.sheets if s.name == "A")
    assert [p.ref for p in a_sheet.pairs()] == ["1.1", "1.2", "1.3"]
    assert [s.heading for s in a_sheet.sections] == [None]  # zero sections
    assert book.skipped == {"unanswered": 2}


def test_non_prose_answers_are_skipped_and_no_figure_reaches_a_warning(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Qs"
    ws["A1"], ws["B1"], ws["C1"] = "Ref", "Question", "Response"
    rows = [("1.1", "Total implementation fee?", "$1,250,000"),
            ("1.2", "Blended discount percentage?", "12%"),
            ("1.3", "Number of consultants?", "40"),
            ("1.4", "Do you subcontract?", "Yes"),
            ("1.5", "Offshore delivery?", "N/A"),
            ("1.6", "Describe your escalation path for severity-one incidents.",
             "A named executive sponsor is engaged within one hour of the call.")]
    for i, (ref, q, a) in enumerate(rows, start=2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = ref, q, a
    fee = wb.create_sheet("Pricing")
    fee["A1"], fee["B1"] = "Item", "Fee"
    fee["A2"], fee["B2"] = "Discovery and planning phase", 40000
    fee["A3"], fee["B3"] = "Build and configuration phase", 120000
    path = tmp_path / "nonprose.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert [p.row for p in book.pairs()] == [7]
    assert book.skipped == {"non-prose answers": 7, "pricing-shaped sheets": 1}
    assert [w for w in book.warnings if w.startswith("Qs!")] == [
        f"Qs!row {r}: non-prose answer skipped" for r in (2, 3, 4, 5, 6)]
    assert "Pricing: pricing-shaped sheet skipped" in book.warnings
    assert not any(tok in w for w in book.warnings
                   for tok in ("1,250", "12%", "40000", "120000"))
    assert all(_shaped(line) for line in book.warnings)


def test_formula_answers_use_the_cached_value_or_warn(tmp_path):
    def build(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "F"
        ws["A1"], ws["B1"], ws["C1"] = "Ref", "Question", "Response"
        ws["A2"], ws["B2"], ws["C2"] = "1.1", "Describe your data migration approach.", "=B2"
        ws["A3"], ws["B3"], ws["C3"] = ("1.2", "Describe your testing approach.",
                                        "Three test cycles, then UAT with agreed exit criteria.")
        wb.save(path)
        return path

    # openpyxl writes no cached value: the formula answer is skipped by name
    plain = build(tmp_path / "plain.xlsx")
    book = pair_response_workbook(plain)
    assert [p.row for p in book.pairs()] == [3]
    assert book.warnings == ["F!row 2: formula answer without cached value skipped"]

    # the cache Excel would have saved, spliced into the sheet XML: the
    # answer IS the cached text
    cached = build(tmp_path / "cached.xlsx")
    _splice_member(
        cached, "xl/worksheets/sheet1.xml",
        b'<c r="C2"><f>B2</f><v></v></c>',
        b'<c r="C2" t="str"><f>B2</f><v>Our migration approach is phased and rehearsed.</v></c>')
    book = pair_response_workbook(cached)
    assert [(p.row, p.answer) for p in book.pairs()][0] == (
        2, "Our migration approach is phased and rehearsed.")
    assert book.warnings == []


def test_hidden_sheets_rows_and_answer_columns_are_skipped_by_the_pairer(tmp_path):
    # hidden-twin: "Vendor Questions" has Ref|Question and no Response
    # column — Layer 2's fallback puts the answer in the HIDDEN column C;
    # "Internal Notes" is a hidden sheet. Nothing hidden is paired, and no
    # directive text reaches a warning.
    book = pair_response_workbook(FIXTURES / "hidden-twin.xlsx")
    assert book.pairs() == []
    assert book.warnings == [
        "Vendor Questions: columns inferred, no header row found",
        "Vendor Questions: hidden answer column, sheet skipped",
        "Internal Notes: hidden sheet skipped",
    ]
    assert not any("rubric" in w.lower() or "compliant" in w.lower()
                   for w in book.warnings)
    # a hidden ROW on a visible, well-formed sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws["A1"], ws["B1"], ws["C1"] = "Ref", "Question", "Response"
    ws["A2"], ws["B2"], ws["C2"] = ("1.1", "Describe your approach.",
                                    "A governed approach with named owners per phase.")
    ws["A3"], ws["B3"], ws["C3"] = ("1.2", "Describe your team.",
                                    "SUPERSEDED DRAFT — do not use this wording.")
    ws.row_dimensions[3].hidden = True
    path = tmp_path / "hiddenrow.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert [p.row for p in book.pairs()] == [2]
    assert book.warnings == ["Questions!row 3: hidden row skipped"]
    assert book.skipped == {"hidden rows": 1}


def test_orphan_answers_and_banner_text_never_reach_a_warning(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws["A1"], ws["B1"], ws["C1"] = "Ref", "Question", "Response"
    ws["A2"], ws["B2"], ws["C2"] = ("1.1", "Describe your approach.",
                                    "A governed approach with named owners per phase.")
    ws["C5"] = "CONFIDENTIAL BANNER TEXT — do not distribute"
    path = tmp_path / "banner.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.warnings == ["Questions!row 5: orphan answer"]
    assert all("CONFIDENTIAL" not in w for w in book.warnings)
    assert all(_shaped(line) for line in book.warnings)


def test_sheet_without_question_and_answer_columns_is_skipped_naming_the_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "3. Pricing"
    ws["A1"], ws["B1"] = "Phase", "Hours"
    ws["A2"], ws["B2"] = "Discover", 400
    ws["A3"], ws["B3"] = "Build", 1200
    path = tmp_path / "grid.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.pairs() == []
    assert book.warnings == ["3. Pricing: no question/answer columns, sheet skipped"]
    assert book.skipped == {"sheets without columns": 1}


def test_a_junk_or_missing_workbook_is_one_typed_refusal(tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip at all")
    try:
        pair_response_workbook(junk)
    except ResponseWorkbookError as exc:
        assert "junk.xlsx" in str(exc)
    else:
        raise AssertionError("a junk container must refuse by type")
    try:
        pair_response_workbook(tmp_path / "absent.xlsx")
    except ResponseWorkbookError as exc:
        assert "absent.xlsx" in str(exc)
    else:
        raise AssertionError("a missing file must refuse by type")


def test_pairing_is_byte_deterministic_across_directories(tmp_path):
    one = pair_response_workbook(fill_demo_twin(_in(tmp_path, "d1")))
    two = pair_response_workbook(fill_demo_twin(_in(tmp_path, "d2")))
    assert json.dumps(one.to_dict(), sort_keys=True) == json.dumps(two.to_dict(), sort_keys=True)
    assert one.source_sha256 == two.source_sha256
    # writers omit: a ref-less pair carries no `ref` key at all
    refless = pair_response_workbook(_refless(tmp_path)).to_dict()
    assert "ref" not in refless["sheets"][0]["sections"][0]["pairs"][0]


def _in(tmp_path, name):
    d = tmp_path / name
    d.mkdir()
    return d / "x.xlsx"


def _refless(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws["A1"], ws["B1"] = "Question", "Response"
    ws["A2"], ws["B2"] = ("Describe your implementation methodology.",
                          "A phased approach with governance gates at each milestone.")
    path = tmp_path / "refless-one.xlsx"
    wb.save(path)
    return path


def test_a_sheet_with_nothing_answered_warns_every_question_row(tmp_path):
    """1.1.0 (P33b, B154 §4): no answered row means no leaf vote. Under
    1.0.0 Layer 2's default depth made every shallow ref row a silent
    heading, so a workbook nobody answered came back with zero pairs AND
    zero addresses — the in-engine reader's refusal would have had nothing
    to say."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Blank"
    ws.append(["Ref", "Question", "Response"])
    ws.append(["1", "Company background", None])
    ws.append(["1.1", "Describe your company history.", None])
    ws.append(["1.2", "Describe your methodology.", None])
    path = tmp_path / "blank.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.pairs() == [] and book.sheets == []
    assert book.skipped == {"unanswered": 3}
    assert book.warnings == ["Blank!row 2: unanswered", "Blank!row 3: unanswered",
                             "Blank!row 4: unanswered"]
    assert all(_shaped(line) for line in book.warnings)


def test_filled_twin_is_byte_frozen_like_the_committed_twins(tmp_path):
    """P33b (B154 §4): private CI 37396189927 went red on the P33a
    determinism test when two fills straddled a second — openpyxl stamps
    dcterms:modified and every zip entry from the wall clock. The filled
    copy now goes through the twins' `_freeze` (EC-8): pinned entry dates,
    pinned core properties, so `source_sha256` is a property of the
    CONTENT and the two determinism tests above cannot flake."""
    import zipfile

    from tests.fixtures.twins import _PINNED_DATE

    path = fill_demo_twin(tmp_path / "frozen.xlsx")
    with zipfile.ZipFile(path) as zf:
        assert {info.date_time for info in zf.infolist()} == {_PINNED_DATE}
        core = zf.read("docProps/core.xml").decode("utf-8")
    assert "<dcterms:modified" in core and ">2026-01-01T00:00:00Z<" in core
    assert "2026-10-" not in core
