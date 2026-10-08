"""P34b (B157): template residue. Three blank buyer templates paired to
1,347, 42 and 3 "answers" under 1.1.0 (B155 §4) — a compliance-code column
taken as the answer column, placeholder text repeated down the response
column, the buyer's own directive left standing in a cell. 1.2.0 chooses
the answer column by content, skips an answer value that stands in three
or more rows, and skips a directive, each by address. Expectations are
hand-derived from the in-test workbooks, never pasted from a run.
"""

import re

from openpyxl import Workbook

from engine.kb.read import read_source
from engine.structure import PARSER_VERSION, parse_workbook
from engine.structure.response import RESPONSE_PARSER_VERSION, pair_response_workbook

_ROW_LINE = re.compile(r"^.+!row \d+: [a-z][a-z -]*[a-z]( \(.+\))?$")
_SHEET_LINE = re.compile(r"^.+: [a-z][a-z/, -]*[a-z]( \([A-Z]+\))?$")

CODES = ["Comply", "Partial", "Exception"]


def _shaped(line: str) -> bool:
    return bool(_ROW_LINE.match(line) or _SHEET_LINE.match(line))


def _prose(i: int) -> str:
    return (f"Our team meets requirement {i} through a documented control that "
            f"is reviewed quarterly by the engagement lead.")


def _matrix(path, *, rows: int = 24, prose: bool = True, comments: bool = True,
            title: str = "Matrix"):
    """A requirements matrix: Ref | Requirement | Compliance Response |
    Vendor Response | Comments. The compliance column cycles three codes;
    the vendor column carries prose (or nothing); comments stay empty."""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    header = ["Ref", "Requirement", "Compliance Response", "Vendor Response"]
    if comments:
        header.append("Comments")
    ws.append(header)
    for i in range(1, rows + 1):
        row = [f"1.{i}", f"Requirement {i}: describe the control for area {i}.",
               CODES[i % 3], _prose(i) if prose else None]
        ws.append(row)
    wb.save(path)
    return path


def test_a_low_distinct_code_column_is_never_the_answer_column(tmp_path):
    path = _matrix(tmp_path / "matrix.xlsx")
    book = pair_response_workbook(path)
    assert book.parser_version == RESPONSE_PARSER_VERSION == "1.2.0"
    assert [p.answer for p in book.pairs()] == [_prose(i) for i in range(1, 25)]
    assert [p.ref for p in book.pairs()] == [f"1.{i}" for i in range(1, 25)]
    assert book.skipped == {"coded answer columns": 1}
    assert book.warnings == ["Matrix: coded answer column skipped (C)"]
    assert all(_shaped(line) for line in book.warnings)
    # the vacancy parser is untouched: same workbook, same pinned version
    assert parse_workbook(path).parser_version == PARSER_VERSION


def test_a_blank_matrix_with_prefilled_codes_pairs_to_zero(tmp_path):
    # the buyer filled the code column before issuing; the vendor column and
    # the comments column are both empty, so the tie keeps the leftmost (D)
    book = pair_response_workbook(_matrix(tmp_path / "blank.xlsx", prose=False))
    assert book.pairs() == [] and book.sheets == []
    assert book.skipped == {"coded answer columns": 1, "unanswered": 24}
    assert book.warnings[0] == "Matrix: coded answer column skipped (C)"
    assert book.warnings[1:] == [f"Matrix!row {r}: unanswered" for r in range(2, 26)]


def test_every_answer_column_coded_skips_the_sheet_by_name(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Codes"
    ws.append(["Ref", "Requirement", "Response"])
    for i in range(1, 25):
        ws.append([f"1.{i}", f"Requirement {i}: describe the control.", CODES[i % 3]])
    path = tmp_path / "codes.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.pairs() == []
    assert book.skipped == {"coded answer columns": 1, "sheets without columns": 1}
    assert book.warnings == ["Codes: coded answer column skipped (C)",
                             "Codes: every answer column coded, sheet skipped"]
    assert all(_shaped(line) for line in book.warnings)


def test_below_the_row_floor_a_column_is_chosen_and_its_repeats_skipped_by_row(tmp_path):
    # ten rows of two codes: under the twenty-row floor the column is not
    # coded, so it IS the answer column — and every value then stands in
    # five rows, so the repeated-value rule names each row
    wb = Workbook()
    ws = wb.active
    ws.title = "Short"
    ws.append(["Ref", "Requirement", "Response"])
    for i in range(1, 11):
        ws.append([f"1.{i}", f"Requirement {i}: describe the control.",
                   ["Comply", "Exception"][i % 2]])
    path = tmp_path / "short.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert book.pairs() == [] and book.sheets == []
    assert book.skipped == {"repeated answer values": 10}
    assert book.warnings == [f"Short!row {r}: repeated answer value skipped"
                             for r in range(2, 12)]


def test_the_longest_median_text_column_wins_over_an_earlier_shorter_one(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Two"
    ws.append(["Ref", "Question", "Response", "Comments"])
    short = ["Supported in release four.", "Configured per tenant.",
             "Delivered by the data team.", "Available on request.",
             "Included in the base scope."]
    long = [f"Answer {i}: a governed, repeatable approach with named owners, "
            f"weekly checkpoints and a documented exit for stage {i}." for i in range(5)]
    for i in range(5):
        ws.append([f"1.{i + 1}", f"Describe stage {i + 1}.", short[i], long[i]])
    path = tmp_path / "two.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert [p.answer for p in book.pairs()] == long
    assert book.warnings == [] and book.skipped == {}


def _two_part(path, placeholder_rows: int):
    """Two sheets, Ref | Question | Response; the first `placeholder_rows`
    rows of each sheet hold the same placeholder (one with odd spacing and
    case), the last row of each a real answer."""
    wb = Workbook()
    first = wb.active
    first.title = "Part A"
    second = wb.create_sheet("Part B")
    real = {"Part A": "A phased approach with governance gates at each milestone.",
            "Part B": "Hypercare for eight weeks, then a tiered service desk."}
    for ws in (first, second):
        ws.append(["Ref", "Question", "Response"])
        for i in range(1, placeholder_rows + 1):
            text = ("To be completed by the respondent." if i % 2
                    else "to  be completed by the RESPONDENT.")
            ws.append([f"1.{i}", f"Describe item {i}.", text])
        ws.append([f"1.{placeholder_rows + 1}", "Describe your support model.", real[ws.title]])
    wb.save(path)
    return path, real


def test_an_answer_value_in_three_or_more_rows_is_skipped_by_address(tmp_path):
    path, real = _two_part(tmp_path / "placeholders.xlsx", placeholder_rows=2)
    book = pair_response_workbook(path)
    assert [(s.name, p.row, p.answer) for s in book.sheets for p in s.pairs()] == [
        ("Part A", 4, real["Part A"]), ("Part B", 4, real["Part B"])]
    assert book.skipped == {"repeated answer values": 4}
    # workbook-wide, so its lines follow every per-sheet line
    assert book.warnings == ["Part A!row 2: repeated answer value skipped",
                             "Part A!row 3: repeated answer value skipped",
                             "Part B!row 2: repeated answer value skipped",
                             "Part B!row 3: repeated answer value skipped"]
    assert all(_shaped(line) for line in book.warnings)


def test_two_identical_answers_still_pair(tmp_path):
    path, _real = _two_part(tmp_path / "twice.xlsx", placeholder_rows=1)
    book = pair_response_workbook(path)
    assert len(book.pairs()) == 4
    assert book.skipped == {} and book.warnings == []


def test_a_directive_left_standing_in_a_response_cell_is_skipped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Q"
    ws.append(["Ref", "Question", "Response"])
    rows = [
        ("1.1", "Describe your approach.",
         "Please describe your approach to data migration in 200 words."),
        ("1.2", "Provide references.",
         "Provide three references from comparable engagements."),
        ("1.3", "Attach your SOC report.",
         "Do not insert here. Attach as a separate appendix."),
        ("1.4", "Describe your support model.",
         "Our approach provides a tiered service desk with named leads."),
        ("1.5", "List your references.",
         "Listed below are three references from comparable public-sector engagements."),
        ("1.6", "Describe your platform.",
         "Enterprise clients run the platform on a dedicated tenant."),
    ]
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "directives.xlsx"
    wb.save(path)
    book = pair_response_workbook(path)
    assert [p.row for p in book.pairs()] == [5, 6, 7]
    assert book.skipped == {"directive answers": 3}
    assert book.warnings == ["Q!row 2: directive answer skipped",
                             "Q!row 3: directive answer skipped",
                             "Q!row 4: directive answer skipped"]
    assert all(_shaped(line) for line in book.warnings)


def test_the_residue_skips_do_not_degrade_the_in_engine_read(tmp_path):
    # a rule is not lost content (B153 §3e): the ingest route carries the
    # warning and leaves the cards un-flagged
    source = read_source(_matrix(tmp_path / "matrix.xlsx"))
    assert source.degraded is False
    assert "Matrix: coded answer column skipped (C)" in source.warnings
    assert source.doc_kind == "past_response"
    assert "Comply" not in source.text and "Exception" not in source.text
