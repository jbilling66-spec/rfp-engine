"""Filled copies of the committed demo twin, built under tmp_path at test
time (P33a, B151 §3f): the buyer's questionnaire with synthetic answers
typed into its response cells — the shape the owner's firm delivers. No
new committed binary: the twin itself stays pinned to its builder
(`tests/fixtures/test_twins.py`), and every answer here is invented
Northwind vocabulary. The filled copy is saved through the twins'
`_freeze` (EC-8): openpyxl re-stamps dcterms:modified and every zip
entry from the wall clock, so two fills a second apart differed in
bytes — and in `source_sha256` — until P33b (private CI 37396189927).
"""

from pathlib import Path

from openpyxl import load_workbook

from tests.fixtures.twins import _freeze

FIXTURES = Path(__file__).resolve().parent
DEMO_TWIN = FIXTURES / "demo-twin.xlsx"


def demo_rows() -> list[tuple[str, int, str, str]]:
    """Every (sheet, row, ref, question) of the committed demo twin's
    question sheets — the rows a filled copy answers. Read from the twin
    the builder pins, in sheet-then-row order."""
    wb = load_workbook(DEMO_TWIN)
    rows = []
    for ws in wb.worksheets:
        if "instruction" in ws.title.lower():
            continue
        for r in range(2, ws.max_row + 1):
            ref, question = ws[f"A{r}"].value, ws[f"B{r}"].value
            if ref and question:
                rows.append((ws.title, r, str(ref), str(question)))
    return rows


def default_answer(ref: str, question: str) -> str:
    """Deterministic prose, always longer than the 60-character label cap
    (the short-answer traps pass their own `answers`)."""
    return (f"For {ref} our delivery team follows a governed, repeatable "
            f"approach refined across prior ERP programs, with named owners "
            f"and weekly checkpoints.")


def fill_demo_twin(dst: Path, *, rows: str = "all", answers=None) -> Path:
    """Write an answer into each response cell of a copy of the demo twin.
    `rows="first"` answers only the first question of every sheet (the
    half-filled shape); `answers(ref, question)` overrides the prose."""
    make = answers or default_answer
    wb = load_workbook(DEMO_TWIN)
    for ws in wb.worksheets:
        if "instruction" in ws.title.lower():
            continue
        for r in range(2, ws.max_row + 1):
            ref, question = ws[f"A{r}"].value, ws[f"B{r}"].value
            if not (ref and question):
                continue
            if rows == "first" and r != 2:
                continue
            ws[f"C{r}"] = make(str(ref), str(question))
    _freeze(wb, dst)  # EC-8: byte-stable across seconds and directories
    return dst
