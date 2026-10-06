"""P33a step 4 (B151 §3d): the `kb pair` door — a completed response
workbook in, the markdown `kb ingest` reads out. Zero spend; counts on
stdout, warnings (addresses only) on stderr; refuses a non-xlsx suffix, an
existing output, a junk container, and an empty result.
"""

import argparse

from openpyxl import Workbook

from engine.cli.main import build_parser, main
from tests.fixtures.filled_twins import fill_demo_twin


def test_kb_pair_writes_markdown_and_reports_counts(tmp_path, capsys):
    src = fill_demo_twin(tmp_path / "filled.xlsx", rows="first")
    out = tmp_path / "out" / "filled.md"
    assert main(["kb", "pair", "--file", str(src), "--out", str(out)]) == 0
    captured = capsys.readouterr()
    assert "8 pairs across 8 sheets" in captured.out
    assert "unanswered 11" in captured.out and "instructions sheets 1" in captured.out
    assert f"wrote {out}" in captured.out
    assert captured.err.count("warning: ") == 11
    assert "warning: 1. Implementation Methodology!row 3: unanswered" in captured.err
    text = out.read_text(encoding="utf-8")
    assert text.startswith("<!-- response-workbook sha256:")
    assert text.count("\n### ") == 8
    assert "filled.xlsx" not in text  # the filename never rides the markdown


def test_kb_pair_refuses_to_overwrite(tmp_path, capsys):
    src = fill_demo_twin(tmp_path / "filled.xlsx")
    out = tmp_path / "existing.md"
    out.write_text("keep me\n", encoding="utf-8")
    assert main(["kb", "pair", "--file", str(src), "--out", str(out)]) == 1
    assert "REFUSED" in capsys.readouterr().err
    assert out.read_text(encoding="utf-8") == "keep me\n"


def test_kb_pair_refuses_non_xlsx_and_bad_containers_typed(tmp_path, capsys):
    docx = tmp_path / "response.docx"
    docx.write_bytes(b"PK\x03\x04 not really")
    out = tmp_path / "a.md"
    assert main(["kb", "pair", "--file", str(docx), "--out", str(out)]) == 1
    err = capsys.readouterr().err
    assert "REFUSED" in err and "'.docx'" in err
    assert not out.exists()
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip at all")
    assert main(["kb", "pair", "--file", str(junk), "--out", str(out)]) == 1
    err = capsys.readouterr().err
    assert "REFUSED" in err and "junk.xlsx" in err and "Traceback" not in err
    assert not out.exists()


def test_kb_pair_exits_nonzero_on_zero_pairs(tmp_path, capsys):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["A1"], ws["B1"] = "Phase", "Hours"
    ws["A2"], ws["B2"] = "Discover", 400
    src = tmp_path / "grid.xlsx"
    wb.save(src)
    out = tmp_path / "grid.md"
    assert main(["kb", "pair", "--file", str(src), "--out", str(out)]) == 1
    captured = capsys.readouterr()
    assert "0 pairs across 0 sheets" in captured.out
    assert "nothing to ingest" in captured.err
    assert "warning: Pricing: no question/answer columns, sheet skipped" in captured.err
    assert not out.exists()


def test_kb_pair_takes_no_kb_root():
    parser = build_parser()
    kb = next(a for a in parser._actions if isinstance(a, argparse._SubParsersAction)
              ).choices["kb"]
    sub = next(a for a in kb._actions if isinstance(a, argparse._SubParsersAction))
    pair = sub.choices["pair"]
    flags = {opt for action in pair._actions for opt in action.option_strings}
    assert "--kb" not in flags
    assert {"--file", "--out"} <= flags
