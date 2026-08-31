"""C9: the docling adapter on the intake path — identity stamped, [page N]
markers preserved, sidecar and grids carried, per-document failure degrades
to legacy without refusing, and the un-backed path is honest about itself."""

import pytest

from engine.intake.extract import ExtractedDoc, extract
from tests.extraction.fakes import FakeExtractionBackend, simple_view
from tests.fixtures.intake_twins import build_minimal_pdf


def _pdf(tmp_path, name="rfp.pdf"):
    path = tmp_path / name
    build_minimal_pdf(path, [["Deadline is March 3, 2027."], ["Second page body."]])
    return path


def test_docling_pdf_keeps_page_markers_and_dates(tmp_path):
    path = _pdf(tmp_path)
    fake = FakeExtractionBackend({
        "rfp.pdf": simple_view(
            "md export", pages=2,
            page_texts=["Deadline is March 3, 2027.", "Second page body."],
        )
    })
    doc = extract(path, backend=fake)
    assert doc.extractor == "docling"
    assert doc.extraction_degraded is False
    assert "[page 1]" in doc.text and "[page 2]" in doc.text
    # The deadline scan works exactly as on the legacy path.
    assert doc.date_candidates[0]["date"] == "2027-03-03"
    assert doc.date_candidates[0]["location"] == "rfp.pdf p1"
    assert doc.extraction_fingerprint.startswith("ext_")


def test_docling_docx_carries_grids_and_sidecar(tmp_path):
    import docx as pydocx

    path = tmp_path / "form.docx"
    d = pydocx.Document()
    d.add_paragraph("Response form")
    d.save(path)
    fake = FakeExtractionBackend({
        "form.docx": simple_view(
            "# Response form\n\nbody",
            grids=[{"grid": [["Item", "Fee"]], "merges": [[[0, 0], [0, 1]]]}],
        ) | {"sidecar": {"fills": {"0": {"FFFF00": [[0, 1]]}}, "comment_texts": []}}
    })
    doc = extract(path, backend=fake)
    assert doc.grids == [{"grid": [["Item", "Fee"]], "merges": [[[0, 0], [0, 1]]]}]
    assert doc.sidecar["fills"]["0"] == {"FFFF00": [[0, 1]]}
    assert doc.text.startswith("# Response form")


def test_empty_pdf_page_warns_like_legacy(tmp_path):
    path = _pdf(tmp_path)
    fake = FakeExtractionBackend({
        "rfp.pdf": simple_view("md", pages=2, page_texts=["Body text here.", ""])
    })
    doc = extract(path, backend=fake)
    assert doc.warnings == ["page 2 produced no text (image-only or empty)"]


def test_partial_success_is_degraded_not_swallowed(tmp_path):
    path = _pdf(tmp_path)
    fake = FakeExtractionBackend({
        "rfp.pdf": simple_view("md", page_texts=["Body text here."],
                               status="partial_success")
    })
    doc = extract(path, backend=fake)
    assert doc.extraction_degraded is True
    assert doc.extraction_flags == ["partial_extraction"]
    assert doc.extractor == "docling"  # degraded, but still the docling read


def test_document_failure_falls_back_to_legacy_degraded(tmp_path):
    path = _pdf(tmp_path)
    fake = FakeExtractionBackend({})  # scripts nothing -> ExtractionFailed
    doc = extract(path, backend=fake)
    assert doc.extractor == "pypdf"
    assert doc.extraction_degraded is True
    assert doc.extraction_flags == ["docling_fallback"]
    assert any("legacy fallback" in w for w in doc.warnings)
    assert "Deadline is March 3, 2027." in doc.text  # legacy read succeeded


def test_no_backend_pdf_is_honest_about_legacy(tmp_path):
    doc = extract(_pdf(tmp_path))
    assert doc.extractor == "pypdf"
    assert doc.extraction_degraded is True
    assert doc.extraction_flags == ["legacy_extractor"]
    assert doc.extraction_fingerprint.startswith("ext_")


def test_no_backend_xlsx_is_primary_not_degraded(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Requirement text"
    wb.save(path)
    doc = extract(path)
    assert doc.extractor == "openpyxl"
    assert doc.extraction_degraded is False
    assert doc.extraction_flags == []


def test_stack_identities_differ_between_paths(tmp_path):
    # Two stacks, two fingerprints — the C12 seam property at doc level.
    legacy = extract(_pdf(tmp_path))
    fake = FakeExtractionBackend({
        "rfp.pdf": simple_view("md", page_texts=["Body text here."])
    })
    adopted = extract(_pdf(tmp_path), backend=fake)
    assert legacy.extraction_fingerprint != adopted.extraction_fingerprint
