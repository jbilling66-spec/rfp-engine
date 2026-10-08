"""Layer 3 for COMPLETED response workbooks (P33a, B151): the buyer's
questionnaire with the firm's answers typed into the response column —
pair each question with the answer beside it, so the knowledge base can be
distilled from what the firm actually sent.

The vacancy parser (classify.py, `parse_workbook`) answers "where can an
answer GO" and is byte-pinned; this module answers "what was SAID" and
carries its own version constant, the DOCX_PARSER_VERSION precedent.
Layer 1 (facts.py) is reused as it is. Layer 2 (conventions.py) is a
named FALLBACK only: on a filled sheet its votes misread the file — the
writable-fill vote dies, the question column ties with the answer column
and flips when a question is short and its answer long, every short
question beside a short answer is a "label row", a filled ref-less
`Question | Response` sheet is a "grid" (B151 §3a). So the header scan,
the leaf depth and the row rules live here, and `sc.label_rows`,
`sc.kind` and `answer_cells` are never read.

Warnings carry addresses and kinds, never a cell's text (P1-23's law):
`"<sheet>!row N: <kind>"` for a row, `"<sheet>: <kind>"` for a sheet.
"""

from __future__ import annotations

import dataclasses
import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile

from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from engine.structure.classify import _field_type
from engine.structure.conventions import (
    depth,
    is_instructions,
    learn_conventions,
    merged_member_cols_by_row,
    row_ref,
)
from engine.structure.facts import CellFact, SheetFacts, collect_workbook_facts
from engine.structure.zipguard import ZipGuardError

# Bumped whenever the pairing rules change. Independent of PARSER_VERSION
# (the vacancy parser's pinned bytes never move for this module).
RESPONSE_PARSER_VERSION = "1.2.0"  # 1.2.0 (P34b): template residue — coded columns, repeated values, directives

HEADER_SCAN_ROWS = 10  # the header is found in the first rows, not voted

# P34b (B157): template residue. A response-vocabulary column whose filled
# cells are a handful of distinct values over many rows is the buyer's
# compliance-code column, never the answer column; an answer value repeated
# across rows is a placeholder, a code or a cross-reference; a cell still
# holding the buyer's own directive was never answered.
CODE_COLUMN_MIN_ROWS = 20
CODE_COLUMN_MAX_DISTINCT = 8
REPEATED_ANSWER_MIN_ROWS = 3

_QUESTION_LABEL = re.compile(
    r"\b(question|requirement|criteri\w*|description|item)\b", re.IGNORECASE)
_ANSWER_LABEL = re.compile(
    r"\b(response|answer|reply|comments?)\b", re.IGNORECASE)
# A later row made only of header words is the header repeated (a long
# sheet re-labelling itself), never a pair.
_HEADER_WORD = re.compile(
    r"^(ref\.?|#|no\.?|id|item|section|question|requirement|criteri\w*|"
    r"description|response|answer|reply|comments?)$", re.IGNORECASE)
# An answer that is a number, a currency or a percentage — never prose,
# never corpus (the P3-1 posture: a fee rides into no card).
_NON_PROSE = re.compile(
    r"^[\s$€£]*[-+]?[\d][\d.,]*\s*(%|k|mm|m|bn)?[\s$€£]*$", re.IGNORECASE)
_BARE_ANSWERS = frozenset({"yes", "no", "y", "n", "n/a", "na", "none", "tbd",
                           "not applicable", "see attached", "see appendix"})
# A response cell still holding the buyer's instruction (P34b): an imperative
# opener ("Provide…", "Please describe…") or the vacancy parser's standing
# phrase. Word-bounded, so "Listed below…" and "Enterprise…" pass.
_DIRECTIVE_OPENER = re.compile(
    r"^\s*(?:please\s+)?(?:provide|describe|insert|enter|list|attach|indicate|"
    r"identify|explain|include|specify|submit|outline|respond)\b", re.IGNORECASE)
_DIRECTIVE_PHRASE = re.compile(r"do not insert here", re.IGNORECASE)


class ResponseWorkbookError(ValueError):
    """The file is not a readable workbook. One typed refusal for the door."""


@dataclass(frozen=True)
class Pair:
    row: int
    question: str  # one line — never reaches a card body
    answer: str
    ref: str | None = None


@dataclass
class ResponseSection:
    heading: str | None
    pairs: list[Pair] = field(default_factory=list)


@dataclass
class ResponseSheet:
    name: str  # BYTE-EXACT, as Layer 1 keeps it
    sections: list[ResponseSection] = field(default_factory=list)

    def pairs(self) -> list[Pair]:
        return [p for s in self.sections for p in s.pairs]


@dataclass
class ResponseWorkbook:
    file: str
    parser_version: str
    source_sha256: str
    sheets: list[ResponseSheet] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped: dict[str, int] = field(default_factory=dict)

    def pairs(self) -> list[Pair]:
        return [p for s in self.sheets for p in s.pairs()]

    def to_dict(self) -> dict:
        """Writers omit: no `None`, sorted skip counts — the shape tests
        compare across directories."""
        raw = dataclasses.asdict(self)
        raw["skipped"] = dict(sorted(raw["skipped"].items()))
        return _omit_none(raw)


def _omit_none(value):
    if isinstance(value, dict):
        return {k: _omit_none(v) for k, v in value.items() if v is not None}
    if isinstance(value, list):
        return [_omit_none(v) for v in value]
    return value


# ---- the door --------------------------------------------------------------

def pair_response_workbook(path: Path) -> ResponseWorkbook:
    path = Path(path)
    if not path.is_file():
        raise ResponseWorkbookError(f"no workbook at {path}")
    try:
        facts = collect_workbook_facts(path)
    except (ZipGuardError, InvalidFileException, BadZipFile, KeyError) as exc:
        raise ResponseWorkbookError(
            f"{path.name}: not a readable workbook ({type(exc).__name__})"
        ) from exc
    conv = learn_conventions(facts)
    book = ResponseWorkbook(
        file=path.name,
        parser_version=RESPONSE_PARSER_VERSION,
        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )
    for sheet in facts.sheets:
        if is_instructions(sheet):
            _skip(book, "instructions sheets")
            continue
        if sheet.hidden:
            _skip(book, "hidden sheets", f"{sheet.name.strip()}: hidden sheet skipped")
            continue
        sc = conv.sheets[sheet.name]
        parsed = _pair_sheet(sheet, sc.ref_col, sc.question_col, sc.answer_col,
                             sc.first_label_row, book)
        if parsed is not None and parsed.pairs():
            book.sheets.append(parsed)
    _drop_repeated_answers(book)  # P34b: a workbook-wide rule, so it runs last
    return book


def _skip(book: ResponseWorkbook, kind: str, warning: str | None = None) -> None:
    book.skipped[kind] = book.skipped.get(kind, 0) + 1
    if warning is not None:
        book.warnings.append(warning)


# ---- one sheet ---------------------------------------------------------------

def _pair_sheet(sheet: SheetFacts, ref_col: int | None, fallback_q: int | None,
                fallback_a: int | None, fallback_header: int | None,
                book: ResponseWorkbook) -> ResponseSheet | None:
    name = sheet.name.strip()
    rows = sheet.rows()
    header_row, q_col, a_col, a_label, coded = _find_header(rows, ref_col)
    for col in coded:  # P34b: a code column is named by letter, then never read
        _skip(book, "coded answer columns",
              f"{name}: coded answer column skipped ({get_column_letter(col)})")
    if header_row is not None and a_col is None:
        _skip(book, "sheets without columns",
              f"{name}: every answer column coded, sheet skipped")
        return None
    if header_row is None:
        q_col, a_col = fallback_q, fallback_a
        if q_col is None or a_col is None or q_col == a_col:
            _skip(book, "sheets without columns",
                  f"{name}: no question/answer columns, sheet skipped")
            return None
        # Layer 2's label row is the header here, or there is none
        header_row = fallback_header or 0
        book.warnings.append(f"{name}: columns inferred, no header row found")
    elif _field_type(a_label) != "text":
        _skip(book, "pricing-shaped sheets", f"{name}: pricing-shaped sheet skipped")
        return None
    if a_col in sheet.hidden_cols:
        _skip(book, "hidden answer columns",
              f"{name}: hidden answer column, sheet skipped")
        return None

    members = merged_member_cols_by_row(sheet)

    def answer_fact(row_num: int, row_facts: list[CellFact]) -> CellFact | None:
        if a_col in members.get(row_num, set()):
            return None  # a merged-range member belongs to its anchor
        return next((f for f in row_facts if f.col == a_col), None)

    # Leaf depth from the rows that ARE answered — never a vacancy vote.
    depths: Counter[int] = Counter()
    for row_num, row_facts in rows.items():
        if row_num <= header_row:
            continue
        ref = row_ref(row_facts, ref_col)
        fact = answer_fact(row_num, row_facts)
        if ref and fact is not None and _answer_text(fact) is not None:
            depths[depth(ref)] += 1
    if depths:
        leaf_depth = max(depths.items(), key=lambda kv: (kv[1], kv[0]))[0]
    else:
        # 1.1.0 (P33b, B154 §4): nothing answered means no leaf to vote.
        # Layer 2's default depth made every shallow ref row a silent
        # heading, so a workbook nobody answered refused with no
        # address at all. No leaves -> no headings -> every question
        # row is `unanswered`, said by address.
        leaf_depth = 0

    out = ResponseSheet(name=sheet.name, sections=[ResponseSection(heading=None)])
    non_prose_here = 0
    for row_num, row_facts in sorted(rows.items()):
        if row_num <= header_row:
            continue  # titles and the header itself
        if row_num in sheet.hidden_rows:
            _skip(book, "hidden rows", f"{name}!row {row_num}: hidden row skipped")
            continue
        if _is_repeated_header(row_facts, ref_col):
            _skip(book, "repeated headers",
                  f"{name}!row {row_num}: repeated header row skipped")
            continue
        ref = row_ref(row_facts, ref_col)
        q_fact = next((f for f in row_facts if f.col == q_col), None)
        if q_fact is not None and q_fact.is_formula and q_fact.cached_text is None:
            _skip(book, "formula questions without cache",
                  f"{name}!row {row_num}: formula question without cached value skipped")
            continue
        question = _one_line(q_fact.text) if q_fact is not None else None
        a_fact = answer_fact(row_num, row_facts)
        if a_fact is not None and a_fact.is_formula and a_fact.cached_text is None:
            _skip(book, "formula answers without cache",
                  f"{name}!row {row_num}: formula answer without cached value skipped")
            continue
        answer = _answer_text(a_fact) if a_fact is not None else None

        if ref and depth(ref) < leaf_depth and answer is None:
            out.sections.append(ResponseSection(heading=question or ref))
            continue
        if question and answer:
            if _non_prose(answer):
                non_prose_here += 1
                _skip(book, "non-prose answers",
                      f"{name}!row {row_num}: non-prose answer skipped")
                continue
            if _directive(answer):
                _skip(book, "directive answers",
                      f"{name}!row {row_num}: directive answer skipped")
                continue
            out.sections[-1].pairs.append(
                Pair(row=row_num, ref=ref, question=question, answer=answer))
            continue
        if question and answer is None:
            _skip(book, "unanswered", f"{name}!row {row_num}: unanswered")
            continue
        if answer and not question:
            _skip(book, "orphan answers", f"{name}!row {row_num}: orphan answer")
            continue
        # neither a question nor an answer: furniture, silent (as classify.py)

    out.sections = [s for s in out.sections if s.pairs]
    if not out.sections and non_prose_here:
        _skip(book, "pricing-shaped sheets", f"{name}: pricing-shaped sheet skipped")
        return None
    return out


# ---- the small rules -------------------------------------------------------------

def _find_header(rows: dict[int, list[CellFact]], ref_col: int | None
                 ) -> tuple[int | None, int | None, int | None, str, list[int]]:
    """The first row (before the first ref-carrying row, within the top
    HEADER_SCAN_ROWS) holding one cell that names the question column and a
    DIFFERENT cell that names the answer column. No length cap — a real
    header reads "Vendor Response (max 200 words; do not alter formatting)".
    Several answer-naming cells (P34b, B157): the choice is by CONTENT, not
    position — code columns are set aside (the fifth value lists them) and
    the longest-median text column among the rest wins; none left means the
    sheet has a header and no answer column (`a_col` None)."""
    for row_num, row_facts in sorted(rows.items())[:HEADER_SCAN_ROWS]:
        if row_ref(row_facts, ref_col) is not None:
            break
        labels = [(f.col, f.text.strip()) for f in row_facts
                  if f.text and not f.is_formula]
        a_cols = [(c, t) for c, t in labels if _ANSWER_LABEL.search(t)]
        q_cols = [c for c, t in labels
                  if _QUESTION_LABEL.search(t) and c not in {ac for ac, _ in a_cols}]
        if not a_cols or not q_cols:
            # a cell may name both ("Response to requirement"): the first
            # question-naming cell that is not the chosen answer cell
            q_only = [c for c, t in labels if _QUESTION_LABEL.search(t)]
            a_only = [(c, t) for c, t in labels if _ANSWER_LABEL.search(t)]
            if len(q_only) >= 1 and len(a_only) >= 1:
                pool = [(c, t) for c, t in a_only if c not in q_only] or [a_only[-1]]
                a_col, a_label, coded = _choose_answer_col(pool, rows, row_num)
                if a_col is None:
                    return row_num, None, None, "", coded
                q_col = next((c for c in q_only if c != a_col), None)
                if q_col is not None:
                    return row_num, q_col, a_col, a_label, coded
            continue
        a_col, a_label, coded = _choose_answer_col(a_cols, rows, row_num)
        if a_col is None:
            return row_num, None, None, "", coded
        return row_num, q_cols[0], a_col, a_label, coded
    return None, None, None, "", []


def _is_repeated_header(row_facts: list[CellFact], ref_col: int | None) -> bool:
    texts = [f.text.strip() for f in row_facts if f.text and not f.is_formula]
    if len(texts) < 2 or row_ref(row_facts, ref_col) is not None:
        return False
    return all(_HEADER_WORD.match(t) for t in texts)


def _one_line(text: str | None) -> str | None:
    """A question is ONE line: a cell's newlines would otherwise put
    question text into the card body (the owner's call, B151 §1c)."""
    if text is None:
        return None
    joined = " ".join(text.split())
    return joined or None


def _answer_text(fact: CellFact) -> str | None:
    """The answer keeps its own line breaks; a formula answer is its cached
    value (Layer 1 put the cache in `text`) or nothing."""
    if fact.text is None:
        return None
    if fact.is_formula and fact.cached_text is None:
        return None
    lines = [line.rstrip() for line in fact.text.splitlines()]
    joined = "\n".join(lines).strip()
    return joined or None


def _non_prose(answer: str) -> bool:
    stripped = answer.strip().strip(".").strip().lower()
    return bool(_NON_PROSE.match(answer)) or stripped in _BARE_ANSWERS


# ---- template residue (P34b, B157) ----------------------------------------------

def _normalized(text: str) -> str:
    return " ".join(text.split()).casefold()


def _directive(answer: str) -> bool:
    return bool(_DIRECTIVE_OPENER.match(answer) or _DIRECTIVE_PHRASE.search(answer))


def _column_profile(rows: dict[int, list[CellFact]], header_row: int, col: int
                    ) -> tuple[int, int, int]:
    """(filled cells, distinct values, median length) of one column below the
    header — whitespace-collapsed, case-folded, and nothing is kept."""
    values: list[str] = []
    for row_num, row_facts in rows.items():
        if row_num <= header_row:
            continue
        fact = next((f for f in row_facts if f.col == col), None)
        text = _answer_text(fact) if fact is not None else None
        if text is not None:
            values.append(_normalized(text))
    if not values:
        return 0, 0, 0
    lengths = sorted(len(v) for v in values)
    return len(values), len(set(values)), lengths[len(lengths) // 2]


def _choose_answer_col(candidates: list[tuple[int, str]],
                       rows: dict[int, list[CellFact]], header_row: int
                       ) -> tuple[int | None, str, list[int]]:
    """Among the answer-naming columns, one whose filled cells are at most
    CODE_COLUMN_MAX_DISTINCT distinct values over at least CODE_COLUMN_MIN_ROWS
    rows is the buyer's code column — set aside, returned by number. Of the
    rest the longest median text wins; a tie keeps the leftmost (1.1.0's
    order, so a workbook with one answer column reads exactly as before)."""
    coded: list[int] = []
    kept: list[tuple[int, int, str]] = []
    for col, label in candidates:
        filled, distinct, median = _column_profile(rows, header_row, col)
        if filled >= CODE_COLUMN_MIN_ROWS and distinct <= CODE_COLUMN_MAX_DISTINCT:
            coded.append(col)
        else:
            kept.append((median, col, label))
    if not kept:
        return None, "", coded
    _median, col, label = max(kept, key=lambda k: (k[0], -k[1]))
    return col, label, coded


def _drop_repeated_answers(book: ResponseWorkbook) -> None:
    """An answer value that stands in REPEATED_ANSWER_MIN_ROWS or more rows
    of the workbook is a placeholder, a code or a cross-reference, never the
    firm's prose — skipped by address. Workbook-wide, so it runs after every
    sheet is paired; its warnings follow the per-sheet ones."""
    counts = Counter(_normalized(p.answer) for p in book.pairs())
    repeated = {k for k, n in counts.items() if n >= REPEATED_ANSWER_MIN_ROWS}
    if not repeated:
        return
    kept_sheets: list[ResponseSheet] = []
    for sheet in book.sheets:
        name = sheet.name.strip()
        for section in sheet.sections:
            kept: list[Pair] = []
            for pair in section.pairs:
                if _normalized(pair.answer) in repeated:
                    _skip(book, "repeated answer values",
                          f"{name}!row {pair.row}: repeated answer value skipped")
                else:
                    kept.append(pair)
            section.pairs = kept
        sheet.sections = [s for s in sheet.sections if s.pairs]
        if sheet.sections:
            kept_sheets.append(sheet)
    book.sheets = kept_sheets
